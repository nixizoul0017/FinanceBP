"""Finance workflow acceptance: independent controls, exact bridges and persistent actions."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import json
import sys
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from analysis.metrics import FinancialEngine, Filters
from analysis.comparison import company_budget_comparison, sales_budget_comparison
from analysis.variance import variance_pack
from analysis.historical import historical_variance
from analysis.forecast import forecast_base
from analysis.scenario import forecast_pack
from etl.sources import load_facts
from etl.controls import source_controls
from finance.budgets import FORMAT, normalize_dataset, check_source, overlay, workbook_bytes, parse_workbook
from finance.store import WorkspaceStore, WorkspaceError
from finance.reconciliation import reconciliation, month_snapshot, compare_snapshot
from finance.issues import detect_issues, validate_notes, validate_thresholds, workflow_snapshot
from reporting.service import build_report, validate_report, markdown_report
from reporting.excel import excel_bytes, inspect_excel
from test_stage9 import simple_history, clone

D=Decimal


def facts():
    f=simple_history()
    f.metadata['source_fingerprint']='independent_fixture'
    f.metadata['source_controls']=source_controls({name:frame.to_dict('records') for name,frame in f.tables.items()},'mapped_facts_v1')
    return f


def dataset(grain=(),mode='amounts',months=range(1,13)):
    d={'format':FORMAT,'entity_id':'EDGE9','year':2026,'input_mode':mode,'grain':list(grain),'expense_grain':[],'sales':[],'expenses':[]}
    for m in months:
        row={'month':f'2026-{m:02}-01',**{key:value for key,value in zip(grain,('A','C','R'))}}
        row.update({'budget_net_revenue':'10000.00','budget_net_cost':'6000.00'} if mode=='amounts' else dict(quantity=100,net_unit_price='100',unit_cost='60',refund_amount='500',cost_reversal_amount='300'))
        d['sales'].append(row)
        d['expenses'].append({'month':f'2026-{m:02}-01','amount':'900.00'})
    return d


def publish(store,context='ctx',data=None,name='年度预算'):
    row=store.save_budget(context,name,'original',data or dataset(),'编制员')
    return store.publish_budget(context,row['id'],'确认员','核对模板及年度金额后发布',row['revision'])


class BudgetWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store=WorkspaceStore(Path(self.tmp.name)/'local.sqlite3')
        self.engine=FinancialEngine(facts())

    def test_draft_does_not_change_analysis_and_publish_is_immutable(self):
        row=self.store.save_budget('ctx','年度预算','original',dataset(),'甲')
        self.assertIs(overlay(self.engine,self.store.budgets('ctx',True)),self.engine)
        with self.assertRaises(WorkspaceError):
            self.store.publish_budget('ctx',row['id'],'','',1)
        pub=self.store.publish_budget('ctx',row['id'],'乙','已核对',1)
        changed=overlay(self.engine,self.store.budgets('ctx',True))
        c=company_budget_comparison(changed,'2026-09',version=pub['id'])
        self.assertEqual(c['sales']['totals']['net_revenue']['budget'],D(10000))
        self.assertEqual(c['expenses']['total']['budget'],D(900))
        self.assertEqual(c['operating_profit']['budget'],D(3100))
        self.assertEqual(self.engine.summary('2026-09'),changed.summary('2026-09'))
        with self.assertRaises(WorkspaceError):
            self.store.save_budget('ctx','年度预算','original',dataset(),'甲',identity=pub['id'],expected_revision=pub['revision'])
        self.assertEqual(self.store.budgets('other'),[])

    def test_adjustment_parent_year_and_optimistic_lock(self):
        parent=publish(self.store)
        row=self.store.save_budget('ctx','调整预算','adjustment',dataset(),'甲',parent_id=parent['id'])
        self.store.save_budget('ctx','调整预算','adjustment',dataset(),'甲',identity=row['id'],expected_revision=row['revision'],parent_id=parent['id'])
        with self.assertRaises(WorkspaceError):
            self.store.save_budget('ctx','调整预算','adjustment',dataset(),'甲',identity=row['id'],expected_revision=row['revision'],parent_id=parent['id'])
        other=dataset();other['year']=2027
        for kind in ('sales','expenses'):
            for r in other[kind]:r['month']=r['month'].replace('2026','2027')
        with self.assertRaises(WorkspaceError):
            self.store.save_budget('ctx','另一年度','adjustment',other,'甲',parent_id=parent['id'])
        self.assertEqual([r['action'] for r in reversed(self.store.audit('ctx'))],['create_draft','publish','create_draft','replace_draft'])

    def test_amount_budget_has_no_fabricated_drivers_and_no_fine_grain_budget(self):
        pub=publish(self.store)
        e=overlay(self.engine,[pub])
        self.assertEqual(variance_pack(e,'2026-09',version=pub['id'])['selected']['revenue_bridge']['reason'],'budget_drivers_missing')
        comparison=sales_budget_comparison(e,'2026-09',version=pub['id'],filters=Filters(products='A'))
        self.assertEqual(comparison['status'],'unsupported_dimensions')
        self.assertIsNone(comparison['totals']['net_revenue']['budget'])
        self.assertNotIn('__UNASSIGNED__',e.dimension_values('product_id'))

    def test_driver_budget_bridge_and_company_expense_grain(self):
        pub=publish(self.store,data=dataset(('product_id','channel_id','region_id'),'drivers'))
        e=overlay(self.engine,[pub])
        v=variance_pack(e,'2026-09',version=pub['id'])
        self.assertTrue(v['company_operating_bridge']['reconciled'])
        self.assertEqual(v['company_operating_bridge']['variance'],D(-100))

    def test_amount_budget_annual_forecast_and_missing_cost(self):
        pub=publish(self.store)
        e=overlay(self.engine,[pub])
        result=forecast_pack(forecast_base(e,'2026-09','recent_mean',3,pub['id']))
        self.assertEqual(result['scenarios']['base']['company_budget_comparison']['net_revenue']['budget'],D(120000))
        self.assertEqual(result['scenarios']['base']['company_budget_comparison']['expenses']['budget'],D(10800))
        d=dataset();d['sales'][8]['budget_net_cost']=None
        pub2=publish(self.store,data=d,name='缺成本预算')
        c=company_budget_comparison(overlay(self.engine,[pub2]),'2026-09',version=pub2['id'])
        self.assertEqual(c['sales']['totals']['net_revenue']['budget'],D(10000))
        self.assertIsNone(c['operating_profit']['budget'])

    def test_duplicate_wrong_year_negative_nonfinite_and_zero(self):
        for change in ('duplicate','year','negative','nan','fraction'):
            d=dataset()
            if change=='duplicate':d['sales'].append(deepcopy(d['sales'][0]))
            if change=='year':d['sales'][0]['month']='2025-01-01'
            if change=='negative':d['sales'][0]['budget_net_revenue']='-1'
            if change=='nan':d['sales'][0]['budget_net_revenue']='NaN'
            if change=='fraction':d['sales'][0]['budget_net_revenue']='1.001'
            with self.subTest(change=change),self.assertRaises(WorkspaceError):normalize_dataset(d)
        d=dataset();d['sales'][0]['budget_net_revenue']=0
        self.assertEqual(normalize_dataset(d)['sales'][0]['budget_net_revenue'],D(0))

    def test_source_entity_dimensions_and_new_plan_codes(self):
        d=dataset(('product_id',));d['sales'][0]['product_id']='NEW'
        with self.assertRaises(WorkspaceError):check_source(d,self.engine)
        self.assertTrue(check_source(d,self.engine,True)['warnings'])
        d['entity_id']='OTHER'
        with self.assertRaises(WorkspaceError):check_source(d,self.engine,True)
        d=dataset();d['expense_grain']=['expense_category']
        for row in d['expenses']:row['expense_category']='NEW_FEE'
        with self.assertRaises(WorkspaceError):check_source(d,self.engine)
        self.assertIn('费用类别',check_source(d,self.engine,True)['warnings'][0])

    def test_populated_template_roundtrip_and_typed_dates_amounts(self):
        raw=workbook_bytes(self.engine,2026,dataset=dataset(('product_id',)))
        result=parse_workbook(raw,self.engine)
        self.assertEqual(result,check_source(dataset(('product_id',)),self.engine))
        wb=load_workbook(BytesIO(raw))
        self.assertEqual(wb.sheetnames,['说明','销售预算','费用预算','分类参考'])
        self.assertEqual(wb['销售预算']['C3'].data_type,'n')
        self.assertEqual(wb['销售预算']['A3'].number_format,'yyyy-mm')
        self.assertTrue(wb['销售预算'].data_validations.count)
        self.assertEqual(wb['销售预算'].freeze_panes,'B3')
        wb.close()

    def test_blank_template_partial_row_and_uncached_formula_rejected(self):
        blank=workbook_bytes(self.engine,2026,['product_id'])
        with self.assertRaises(WorkspaceError):parse_workbook(blank,self.engine)
        wb=load_workbook(BytesIO(blank));wb['销售预算']['B3']='A'
        out=BytesIO();wb.save(out)
        with self.assertRaises(WorkspaceError):parse_workbook(out.getvalue(),self.engine)
        wb['销售预算']['C3']='=100*2'
        out=BytesIO();wb.save(out);wb.close()
        with self.assertRaisesRegex(WorkspaceError,'缓存'):parse_workbook(out.getvalue(),self.engine)

    def test_concurrent_publication_allows_only_one_transition(self):
        row=self.store.save_budget('ctx','并发预算','original',dataset(),'甲')
        def attempt(_):
            try:self.store.publish_budget('ctx',row['id'],'乙','核对',1);return True
            except WorkspaceError:return False
        with ThreadPoolExecutor(2) as pool:result=list(pool.map(attempt,range(2)))
        self.assertEqual(sum(result),1)


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.store=WorkspaceStore(Path(self.tmp.name)/'local.sqlite3')
        self.f=facts();self.e=FinancialEngine(self.f)
        # Independent known fixture: 100×100 revenue, 5×100 refunds, 100×60 cost,
        # 5×60 recoveries and a 1000 company bill.
        self.refs=dict(sales_revenue='10000',refund_amount='500',sales_cost='6000',cogs_reversal_amount='300',expenses='1000')

    def controls(self):
        return self.store.save_controls('ctx','2026-09',self.refs,'核对员','独立测试账簿，2026年9月')

    def test_independent_control_prerequisites_and_tieout(self):
        self.assertFalse(reconciliation(self.e,'2026-09')['can_confirm'])
        self.assertTrue(reconciliation(self.e,'2026-09',self.controls())['can_confirm'])
        with self.assertRaises(WorkspaceError):self.store.confirm_month('other',self.e,'2026-09','甲','已核对')
        self.refs['sales_cost']='5999.98'
        self.assertFalse(reconciliation(self.e,'2026-09',self.controls())['can_confirm'])

    def test_source_to_canonical_discrepancy_is_blocked(self):
        self.f.tables['sales'].loc[20,'sales_amount_before_refund']+=D(1)
        result=reconciliation(FinancialEngine(self.f),'2026-09',self.controls())
        self.assertFalse(result['can_confirm'])
        self.assertIn('内部勾稽不一致','；'.join(result['blockers']))

    def test_missing_source_cost_and_incomplete_month_cannot_confirm(self):
        for kind in ('cost','refund','period'):
            f=clone(self.f)
            if kind=='cost':f.tables['sales'].loc[20,'cogs_amount']=None
            if kind=='refund':f.metadata['data_availability']['refunds']=False
            if kind=='period':f.metadata['coverage']['incomplete_months']=['2026-09']
            self.assertFalse(reconciliation(FinancialEngine(f),'2026-09',self.controls())['can_confirm'])

    def test_confirm_change_detection_release_and_reconfirm(self):
        same_second=patch('finance.store.now',return_value='2026-10-16T09:00:00+08:00')
        same_second.start();self.addCleanup(same_second.stop)
        self.controls();identity=self.store.confirm_month('ctx',self.e,'2026-09','甲','完整月核对')
        with self.assertRaises(WorkspaceError):self.store.confirm_month('ctx',self.e,'2026-09','甲','重复确认')
        snapshot=self.store.closes('ctx')[0]['snapshot']
        self.assertFalse(compare_snapshot(snapshot,self.e,'2026-09')['changed'])
        f=clone(self.f);f.tables['sales'].loc[20,'sales_amount_before_refund']+=D(10)
        result=compare_snapshot(snapshot,FinancialEngine(f),'2026-09')
        self.assertTrue(result['changed']);self.assertEqual(result['rows'][0]['kind'],'修改')
        self.store.release_month('ctx',identity,'乙','更正原始记录前解除')
        self.store.confirm_month('ctx',self.e,'2026-09','甲','复核后重新确认')
        self.assertEqual([r['revision'] for r in self.store.closes('ctx')],[1,2])

    def test_budget_refresh_and_other_months_do_not_invalidate_confirmation(self):
        snapshot=month_snapshot(self.e,'2026-09')
        pub=publish(self.store)
        e=overlay(self.e,[pub])
        self.assertFalse(compare_snapshot(snapshot,e,'2026-09')['changed'])
        f=clone(self.f);f.tables['sales'].loc[0,'sales_amount_before_refund']+=D(1)
        f.metadata['source_fingerprint']='newmonthdata'
        self.assertFalse(compare_snapshot(snapshot,FinancialEngine(f),'2026-09')['changed'])

    def test_source_controls_small_and_dynamic_adapters(self):
        for file in ('sample_files.json','multidimensional_files.json'):
            e=FinancialEngine(load_facts(ROOT/'configs'/file))
            rec=reconciliation(e,'2026-09')
            self.assertTrue(all(row['内部勾稽']=='一致' for row in rec['rows']))


class HistoricalAndActionTests(unittest.TestCase):
    def setUp(self):
        self.f=facts();self.e=FinancialEngine(self.f)
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.store=WorkspaceStore(Path(self.tmp.name)/'local.sqlite3')

    def declining(self):
        f=clone(self.f)
        f.tables['sales'].loc[20,'quantity']=80
        f.tables['sales'].loc[20,'sales_amount_before_refund']=D(7200)
        f.tables['sales'].loc[20,'cogs_amount']=D(5600)
        return FinancialEngine(f)

    def test_actual_bridges_reconcile_without_any_budget(self):
        f=clone(self.f)
        for name in ('sales_budget','expense_budget'):f.tables[name]=f.tables[name].iloc[:0];f.metadata['data_availability'][name]=False
        f.tables['sales'].loc[20,'quantity']=120
        f.tables['sales'].loc[20,'sales_amount_before_refund']=D(13200)
        f.tables['sales'].loc[20,'cogs_amount']=D(7800)
        h=historical_variance(FinancialEngine(f),'2026-09',group_by=('product_id','channel_id'))
        components={r['factor']:r['amount'] for r in h['revenue_bridge']['components']}
        self.assertEqual(components['volume'],D(2000));self.assertEqual(components['price'],D(1200))
        for key in ('sales_bridge','revenue_bridge','gross_margin_bridge','operating_bridge'):
            b=h[key];self.assertTrue(b['reconciled']);self.assertAlmostEqual(sum(r['amount'] for r in b['components']),b['variance'])
            self.assertNotIn('budget',b)
        self.assertEqual(h['rankings'][0]['net_revenue']['change_amount'],D(3200))

    def test_history_missing_and_sales_scope_profit_not_fabricated(self):
        self.assertEqual(historical_variance(self.e,'2025-01')['revenue_bridge']['reason'],'base_incomplete')
        h=historical_variance(self.e,'2026-09',Filters(products='A'),mode='yoy')
        self.assertEqual(h['metadata']['base_month'],date(2025,9,1))
        self.assertEqual(h['operating_bridge']['reason'],'sales_scope_unallocated')
        f=clone(self.f);f.tables['sales'].loc[20,'cogs_amount']=None
        h=historical_variance(FinancialEngine(f),'2026-09')
        self.assertEqual(h['revenue_bridge']['status'],'complete')
        self.assertEqual(h['gross_margin_bridge']['reason'],'historic_cost_missing')

    def test_dynamic_cross_period_refund_bridge_and_missing_bucket(self):
        e=FinancialEngine(load_facts(ROOT/'configs/multidimensional_files.json'))
        h=historical_variance(e,'2026-08',group_by=('sales_department_id','salesperson_id'))
        self.assertTrue(h['revenue_bridge']['reconciled'])
        self.assertTrue(h['gross_margin_bridge']['reconciled'])
        self.assertEqual(sum(row['net_revenue']['change_amount'] for row in h['rankings']),h['revenue_bridge']['variance'])

    def test_thresholds_require_both_amount_and_rate(self):
        low=detect_issues(self.declining(),'2026-09',version='V1',thresholds={'amount':100000,'rate':0})
        self.assertEqual(low['candidates'],[])
        high=detect_issues(self.declining(),'2026-09',version='V1')
        self.assertIn('revenue_decline',[r['rule'] for r in high['candidates']])
        self.assertIn('margin_decline',[r['rule'] for r in high['candidates']])
        self.assertEqual(validate_thresholds({'margin_points':'0.025'})['margin_points'],D('.025'))
        self.assertEqual(detect_issues(self.e,'2026-09',version='V1',thresholds={'amount':0,'rate':0,'margin_points':0})['candidates'],[])

    def test_issue_refresh_preserves_actions_and_does_not_auto_resolve(self):
        result=detect_issues(self.declining(),'2026-09',version='V1')
        self.store.scan_issues('ctx','2026-09',result['candidates'],'甲',result['scan_key'])
        row=self.store.issues('ctx')[0]
        notes=dict(owner='业务经理',action='核对促销成交价',due_date='2026-10-20',cause_status='pending')
        self.store.update_issue('ctx',row['id'],notes,'in_progress','乙',row['revision'])
        self.store.scan_issues('ctx','2026-09',[],'甲',result['scan_key'])
        changed=next(r for r in self.store.issues('ctx') if r['id']==row['id'])
        self.assertEqual(changed['status'],'in_progress')
        self.assertEqual(changed['notes']['owner'],'业务经理')
        self.assertFalse(changed['facts']['currently_detected'])
        with self.assertRaises(WorkspaceError):self.store.update_issue('ctx',row['id'],notes,'open','乙',row['revision'])

    def test_scan_scope_isolation_and_notes_validation(self):
        result=detect_issues(self.declining(),'2026-09',version='V1')
        self.store.scan_issues('ctx','2026-09',result['candidates'],'甲',result['scan_key'])
        self.store.scan_issues('ctx','2026-09',[],'甲','different_scope')
        self.assertTrue(all(r['facts']['currently_detected'] for r in self.store.issues('ctx')))
        for notes,status in [({},'in_progress'),({},'resolved'),({'cause_status':'confirmed','cause':'活动'},'open'),({'review_month':'2026-10'},'open')]:
            with self.assertRaises(WorkspaceError):validate_notes(notes,status)
        self.assertEqual(validate_notes({'resolution':'已核实无需处理'},'ignored')['resolution'],'已核实无需处理')

    def test_report_workflow_history_exports_and_stale_identity(self):
        engine=self.declining()
        result=detect_issues(engine,'2026-09',version='V1')
        self.store.scan_issues('ctx','2026-09',result['candidates'],'甲',result['scan_key'])
        work=workflow_snapshot(self.store,'ctx',engine,'2026-09')
        self.assertTrue(all(row.get('scope_label') for row in work['issues']))
        h=historical_variance(engine,'2026-09')
        report=build_report(engine,'2026-09',version='V1',method='recent_mean',workflow=work,historical=h)
        self.assertEqual(validate_report(report)['status'],'passed')
        self.assertIn('对账与人工跟进',markdown_report(report))
        book=inspect_excel(excel_bytes(report));self.assertEqual(len(book.sheetnames),10)
        self.assertIn('人工原因',str([cell.value for row in book['口径与追溯'] for cell in row]))
        broken=deepcopy(report);broken['sources']['workflow']['issues'][0]['notes']['owner']='篡改'
        with self.assertRaises(ValueError):excel_bytes(broken)
        row=self.store.issues('ctx')[0]
        self.store.update_issue('ctx',row['id'],{'resolution':'核实后忽略'},'ignored','乙',row['revision'])
        updated=build_report(engine,'2026-09',version='V1',method='recent_mean',workflow=workflow_snapshot(self.store,'ctx',engine,'2026-09'),historical=h)
        self.assertNotEqual(report['report_id'],updated['report_id'])


if __name__=='__main__':unittest.main()
