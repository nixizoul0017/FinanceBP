"""Real workbook/Stata readers and their financial import boundary."""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch
import uuid

import numpy as np
import pandas as pd
from openpyxl import Workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from analysis import FinancialEngine, Filters
from dashboard.uploads import SelectedFile
from desktop.events import atomic_json
from desktop.paths import AppPaths
from etl.file_tables import list_sheets, read_table
from etl.file_import import build_profile, saved_profiles, save_profile, stage_uploads, suggest_mapping, table_references
from etl.sources import CONTRACT, load_facts
from etl.dimensions import MISSING


def simple_sales():
    return [{'明细编号':'0001','订单号':'0010','销售日期':'2026-09-03','商品':'P1','渠道':'门店','地区':'华北',
             '数量':2,'销售金额':300,'销售成本':100,'员工编号':'007','销售部门':'D1'}]


def write_book(path,tables):
    book=Workbook();book.remove(book.active)
    for name,rows in tables.items():
        sheet=book.create_sheet(name)
        columns=list(rows[0]) if rows else CONTRACT[name]
        sheet.append(columns)
        for row in rows:
            sheet.append([None if row.get(key)==MISSING else float(row[key]) if isinstance(row.get(key),Decimal) else row.get(key) for key in columns])
    book.save(path);book.close()


class BusinessFilesTests(unittest.TestCase):
    def setUp(self):
        temporary=tempfile.TemporaryDirectory();self.addCleanup(temporary.cleanup)
        self.root=Path(temporary.name)
        self.paths=AppPaths(ROOT,self.root/'personal')

    def references(self,path):
        return table_references(stage_uploads([SelectedFile(path.read_bytes(),path.name,str(path))],paths=self.paths))

    def profile(self,path,**options):
        refs=self.references(path)
        selections={}
        for ref in refs:
            role=ref['sheet_name'] or 'sales'
            preview=read_table(ref['path'],sheet_name=ref['sheet_name'])
            selections[role]={'reference':ref,'mapping':suggest_mapping(role,preview.columns,preview.comments)}
        selections.setdefault('refunds',{'empty':True})
        return build_profile('TEST','测试导入','2026-09','2026-09',selections,**options)

    def test_xlsx_sheets_dates_text_codes_and_formatted_numeric_codes(self):
        path=self.root/'中文订单.xlsx';book=Workbook();sheet=book.active;sheet.title='销售明细'
        sheet.append(['编号','日期','员工','金额']);sheet.append(['0001',datetime(2026,9,3),7,135.25]);sheet['C2'].number_format='0000'
        book.create_sheet('退款');book.save(path);book.close()
        self.assertEqual(list_sheets(path),['销售明细','退款'])
        table=read_table(path,sheet_name='销售明细')
        self.assertEqual(table.rows,[{'编号':'0001','日期':'2026-09-03','员工':'0007','金额':135.25}])

    def test_real_legacy_xls_date_and_zero_padding(self):
        table=read_table(ROOT/'tests/fixtures/business_import_legacy.xls',sheet_name='销售明细')
        self.assertEqual(table.rows[0],{'销售明细ID':'00001','销售日期':'2026-09-03','销售金额':135.25,'员工编号':'0007'})

    def test_stata_dates_unicode_labels_missing_and_original_value_codes(self):
        path=self.root/'销售.dta'
        pd.DataFrame({'staff':[7,8],'day':pd.to_datetime(['2026-09-03','2026-09-04']),
                      'amount':[135.25,np.nan]}).to_stata(path,write_index=False,version=118,
                      convert_dates={'day':'td'},variable_labels={'staff':'销售员编号','amount':'销售金额'},
                      value_labels={'staff':{7:'张三',8:'李四'}})
        table=read_table(path)
        self.assertEqual(table.rows[0],{'staff':7,'day':'2026-09-03','amount':135.25})
        self.assertIsNone(table.rows[1]['amount'])
        self.assertEqual(table.comments['staff'],'销售员编号')
        self.assertEqual(suggest_mapping('sales',table.columns,table.comments)['sales_amount_before_refund'],'amount')
        self.assertEqual(len(read_table(path,limit=1).rows),1)

    def test_csv_gb18030_tsv_bom_and_leading_zeros(self):
        for suffix,encoding,delimiter in [('csv','gb18030',','),('tsv','utf-16','\t'),('csv','utf-8-sig',',')]:
            with self.subTest(encoding=encoding):
                path=self.root/('数据.'+suffix)
                path.write_text(delimiter.join(['编号','金额'])+'\n'+delimiter.join(['0007','0.10']),encoding=encoding)
                self.assertEqual(read_table(path).rows,[{'编号':'0007','金额':'0.10'}])

    def test_json_records_remain_compatible_but_config_json_is_distinguished(self):
        path=self.root/'data.json';path.write_text('[{"a":"001"},{"a":"002","b":1}]')
        self.assertEqual(read_table(path).rows[0],{'a':'001','b':None})
        path.write_text('[]');self.assertEqual(read_table(path).rows,[])
        path.write_text('{"profile_version":2}')
        with self.assertRaisesRegex(ValueError,'配置JSON'):read_table(path)

    def test_header_offset_duplicate_and_unheaded_data_are_rejected(self):
        path=self.root/'data.csv';path.write_text('title\n编号,金额\n001,10\n',encoding='utf-8')
        self.assertEqual(read_table(path,header_row=2).rows[0]['编号'],'001')
        for data,reason in [('a,a\n1,2','重复'),('a,,b\n1,2,3','空白'),('a,b\n1,2,3','表头')]:
            path.write_text(data)
            with self.assertRaisesRegex(ValueError,reason):read_table(path)

    def test_errors_corruption_and_size_limits(self):
        path=self.root/'fake.xlsx';path.write_text('not a workbook')
        with self.assertRaisesRegex(ValueError,'损坏'):read_table(path)
        path=self.root/'error.xlsx';book=Workbook();book.active.append(['amount']);book.active.append(['#DIV/0!']);book.save(path);book.close()
        with self.assertRaisesRegex(ValueError,'错误单元格'):read_table(path)
        with patch('etl.file_tables.MAX_FILE_BYTES',1):
            with self.assertRaisesRegex(ValueError,'100MB'):read_table(path)
        path=self.root/'data.csv';path.write_text('a,b\n1,2\n3,4')
        with patch('etl.file_tables.MAX_ROWS',1):
            with self.assertRaisesRegex(ValueError,'100万行'):read_table(path)
        self.assertEqual(len(read_table(path,limit=1).rows),1)

    def test_uploaded_files_are_private_deduplicated_and_saved_profiles_are_stable(self):
        path=self.root/'sales.xlsx';write_book(path,{'sales':simple_sales()})
        uploaded=SelectedFile(path.read_bytes(),'../销售.xlsx',str(path))
        staged=stage_uploads([uploaded,uploaded],paths=self.paths)
        self.assertEqual(len(staged),1)
        self.assertTrue(Path(staged[0]['path']).is_relative_to(self.paths.data/'imports'))
        self.assertEqual(staged[0]['name'],'销售.xlsx')
        profile=self.profile(path);save_profile(profile,paths=self.paths)
        self.assertEqual(saved_profiles(paths=self.paths),[profile])
        write_book(path,{'sales':[{**simple_sales()[0],'销售金额':400}]})
        refreshed=self.profile(path)
        self.assertEqual(profile['profile_id'],refreshed['profile_id'])
        self.assertNotEqual(profile['tables']['sales']['files'],refreshed['tables']['sales']['files'])

    def test_required_mapping_and_duplicate_table_roles_cannot_be_applied(self):
        path=self.root/'sales.xlsx';write_book(path,{'sales':simple_sales()})
        ref=self.references(path)[0]
        with self.assertRaisesRegex(ValueError,'缺少字段映射'):
            build_profile('TEST','导入','2026-09','2026-09',{'sales':{'reference':ref,'mapping':{}}})
        mapping=suggest_mapping('sales',list(simple_sales()[0]))
        with self.assertRaisesRegex(ValueError,'不能同时'):
            build_profile('TEST','导入','2026-09','2026-09',{'sales':{'reference':ref,'mapping':mapping},'refunds':{'reference':ref,'mapping':{}}})

    def test_sales_only_import_dynamic_classification_missing_cost_and_missing_budget(self):
        path=self.root/'sales.xlsx';write_book(path,{'sales':simple_sales()})
        profile=self.profile(path,dimension_fields=['员工编号','销售部门'])
        engine=FinancialEngine(load_facts(profile))
        self.assertEqual(engine.summary('2026-09')['net_revenue'],Decimal(300))
        self.assertEqual(len(engine.dimensions('2026-09',by=[d['id'] for d in profile['dimensions']])),1)
        self.assertFalse(engine.metadata['data_availability']['sales_budget'])
        self.assertIsNone(engine.summary('2026-09')['expenses'])
        profile['tables']['sales']['mapping'].pop('cogs_amount')
        self.assertIsNone(FinancialEngine(load_facts(profile)).summary('2026-09')['net_cost'])

    def test_missing_refunds_and_confirmed_zero_are_different(self):
        path=self.root/'sales.xlsx';write_book(path,{'sales':simple_sales()})
        profile=self.profile(path)
        self.assertEqual(FinancialEngine(load_facts(profile)).summary('2026-09')['net_revenue'],Decimal(300))
        profile['tables'].pop('refunds')
        self.assertIsNone(FinancialEngine(load_facts(profile)).summary('2026-09')['net_revenue'])

    def test_strict_stata_float_money_requires_explicit_rounding(self):
        path=self.root/'sales.dta'
        frame=pd.DataFrame([{'sale_line_id':'001','order_id':'001','revenue_date':'2026-09-03','quantity':1,
                             'sales_amount_before_refund':np.float32(0.1),'cogs_amount':0}])
        frame.to_stata(path,write_index=False,version=118)
        profile=self.profile(path)
        with self.assertRaisesRegex(ValueError,'invalid_amount'):load_facts(profile)
        profile['tables']['sales']['money_rounding']='half_up'
        self.assertEqual(FinancialEngine(load_facts(profile)).summary('2026-09')['net_revenue'],Decimal('0.10'))

    def test_five_table_excel_financial_equivalence_and_refund_link(self):
        original=load_facts(ROOT/'configs/sample_files.json')
        tables={role:frame.to_dict('records') for role,frame in original.tables.items()}
        # Source sample includes a deliberately incomplete October boundary event.
        tables['refunds']=[row for row in tables['refunds'] if str(row['refund_date'])<='2026-09-30']
        path=self.root/'五表.xlsx';write_book(path,tables)
        profile=self.profile(path)
        actual=FinancialEngine(load_facts(profile)).summary('2026-09')
        expected=FinancialEngine(original).summary('2026-09')
        for key in ['net_revenue','net_cost','gross_margin','expenses','operating_profit']:
            self.assertEqual(actual[key],expected[key],key)
        self.assertEqual(actual['net_revenue'],Decimal(3240))
        bad=deepcopy(profile);bad['tables']['refunds']['mapping']['sale_line_id']='refund_id'
        with self.assertRaisesRegex(ValueError,'unknown_sale'):load_facts(bad)

    def test_stata_sales_refunds_expenses_equivalent_to_original_sample(self):
        original=load_facts(ROOT/'configs/sample_files.json');selections={}
        for role,frame in original.tables.items():
            if role=='refunds':frame=frame[frame['refund_date'].astype(str).le('2026-09-30')]
            converted=frame.drop(columns=['entity_id']).copy()
            for field in converted:
                if field in ('revenue_date','refund_date','expense_date','month'):
                    converted[field]=pd.to_datetime(converted[field])
                elif any(isinstance(value,Decimal) for value in converted[field]):
                    converted[field]=converted[field].astype(float)
            path=self.root/(role+'.dta');converted.to_stata(path,write_index=False,version=118)
            ref=self.references(path)[0]
            selections[role]={'reference':ref,'mapping':{field:field for field in converted}}
        profile=build_profile('TEST','dta五表','2026-09','2026-09',selections)
        self.assertEqual(FinancialEngine(load_facts(profile)).summary('2026-09')['net_revenue'],Decimal(3240))

    def test_imported_multi_dimensions_keep_4300_and_historical_refund_department(self):
        original=load_facts(ROOT/'configs/multidimensional_files.json')
        path=self.root/'多维.xlsx'
        write_book(path,{role:original.tables[role].to_dict('records') for role in ('sales','refunds','expenses')})
        profile=self.profile(path,dimension_fields=['salesperson_id','sales_department_id'])
        engine=FinancialEngine(load_facts(profile))
        self.assertEqual(engine.summary('2026-09',filters=Filters(dimensions={'salesperson_id':['1']}))['net_revenue'],Decimal(4300))
        selected=Filters(dimensions={'sales_department_id':['D1']})
        self.assertEqual(engine.period_events('2026-08','2026-08',selected)['refund_amount'].sum(),Decimal(100))

    def test_backup_includes_saved_mapping_but_excludes_business_payload(self):
        import zipfile
        from desktop.workspace import backup, restore
        path=self.root/'sales.xlsx';write_book(path,{'sales':simple_sales()})
        profile=self.profile(path);saved=save_profile(profile,paths=self.paths)
        archive=self.root/'backup.zip';backup(self.paths.data,archive)
        with zipfile.ZipFile(archive) as book:
            self.assertIn('data/profiles/business_imports/'+saved.name,book.namelist())
            self.assertFalse(any('/imports/' in name for name in book.namelist()))
        destination=self.root/'restored';restore(destination,archive)
        self.assertEqual(json.loads((destination/'profiles/business_imports'/saved.name).read_text(encoding='utf-8')),profile)

    def native_upload(self,files,*,multiple=True,types=('xlsx','dta'),limit=100*1024*1024):
        from desktop.host import Host
        request=self.root/(uuid.uuid4().hex+'.json')
        atomic_json(request,{'id':request.stem,'types':list(types),'limit':limit,'multiple':multiple})
        host=Host(self.paths,ROOT/'scripts/windows_app.py',None);host.window=MagicMock()
        host.window.create_file_dialog.return_value=tuple(str(path) for path in files)
        host.upload(request)
        return request,json.loads(request.with_suffix('.result').read_text()),host

    def test_native_multiple_files_and_legacy_single_selection(self):
        a=self.root/'中文.xlsx';a.write_bytes(b'xlsx');b=self.root/'sales.dta';b.write_bytes(b'dta')
        request,result,host=self.native_upload([a,b])
        self.assertEqual(result['state'],'selected')
        self.assertEqual([item['name'] for item in result['files']],['中文.xlsx','sales.dta'])
        self.assertEqual((self.root/(request.stem+'_1.bin')).read_bytes(),b'dta')
        self.assertTrue(host.window.create_file_dialog.call_args.kwargs['allow_multiple'])
        request,result,_=self.native_upload([a],multiple=False)
        self.assertEqual(result['name'],'中文.xlsx');self.assertEqual(request.with_suffix('.bin').read_bytes(),b'xlsx')

    def test_native_cancel_bad_extension_and_partial_error_clean_up(self):
        _,result,_=self.native_upload([]);self.assertEqual(result['state'],'cancelled')
        a=self.root/'sales.xlsx';a.write_bytes(b'xlsx');b=self.root/'bad.exe';b.write_bytes(b'no')
        request,result,_=self.native_upload([a,b])
        self.assertEqual(result['state'],'error')
        self.assertFalse(list(self.root.glob(request.stem+'*.bin')))
        _,result,_=self.native_upload([a],limit=1);self.assertEqual(result['state'],'error')


if __name__=='__main__':unittest.main()
