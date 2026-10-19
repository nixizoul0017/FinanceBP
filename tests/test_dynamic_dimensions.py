"""0.4.0 financial, import, persistence and export contract acceptance."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from analysis import FinancialEngine, Filters
from analysis.comparison import sales_budget_comparison
from analysis.forecast import History, predict, period_summary, forecast_base
from analysis.scenario import forecast_pack
from dashboard.dimension_store import DimensionStore, clean_layout
from dashboard.view_model import detail_tables, safe_export_name, forecast_groups
from etl.dimensions import BUILTINS, MISSING, attach_dimensions, profile_key, recommend, validate_dimensions
from etl.sources import inspect_profile, load_facts, validate_profile
from reporting.service import build_report
from reporting.excel import excel_bytes
from openpyxl import load_workbook


def profile():
    return json.loads((ROOT/'configs/multidimensional_files.json').read_text(encoding='utf-8'))


class DynamicFinancialTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.facts=load_facts(profile())
        cls.engine=FinancialEngine(cls.facts)

    def test_joint_group_totals_rate_and_cross_group_order_deduplication(self):
        engine=self.engine
        grouped=engine.dimensions('2026-09',by=['sales_department_id','salesperson_id'])
        total=engine.summary('2026-09')
        self.assertEqual(sum(grouped['net_revenue']),total['net_revenue'])
        self.assertEqual(sum(grouped['gross_margin']),total['gross_margin'])
        for row in grouped.to_dict('records'):
            self.assertEqual(row['gross_margin_rate'],row['gross_margin']/row['net_revenue'])
        by_product=engine.dimensions('2026-09',by='product_id')
        self.assertGreater(sum(by_product['confirmed_orders']),total['confirmed_orders'])
        reverse=engine.dimensions('2026-09',by=['salesperson_id','sales_department_id'])
        self.assertEqual(list(reverse)[:2],['salesperson_id','sales_department_id'])

    def test_filters_intersection_union_and_company_fees(self):
        selected=Filters(dimensions={'salesperson_id':['1','2'],'sales_department_id':['D2']})
        result=self.engine.summary('2026-09',filters=selected)
        person=self.engine.summary('2026-09',filters=Filters(dimensions={'salesperson_id':['1']}))
        self.assertEqual(result['net_revenue'],person['net_revenue'])
        self.assertIsNone(result['expenses'])
        self.assertIsNone(result['operating_profit'])
        self.assertEqual(self.engine.summary('2026-09')['expenses'],Decimal(1200))
        self.assertEqual(Filters(dimensions={'salesperson_id':['1'],'product_id':['P1']}),Filters(products=['P1'],dimensions={'salesperson_id':['1']}))

    def test_refund_after_transfer_keeps_original_sale_department(self):
        refunds=detail_tables(self.engine,'2026-08')['refunds']
        self.assertEqual(refunds.iloc[0]['sales_department_id'],'D1')
        self.assertEqual(refunds.iloc[0]['salesperson_id'],'1')
        current=self.engine.tables['sales']
        transferred=current[current['sale_line_id'].eq('S202608-1-1')].iloc[0]
        self.assertEqual(transferred['sales_department_id'],'D2')
        events=self.engine.period_events('2026-08','2026-08',Filters(dimensions={'sales_department_id':['D1']}))
        self.assertEqual(events['refund_amount'].sum(),Decimal(100))

    def test_missing_bucket_keeps_every_amount(self):
        grouped=self.engine.dimensions('2026-08',by='salesperson_id')
        missing=grouped[grouped['salesperson_id'].eq(MISSING)]
        self.assertEqual(len(missing),1)
        self.assertGreater(missing.iloc[0]['net_revenue'],0)
        self.assertEqual(sum(grouped['net_revenue']),self.engine.summary('2026-08')['net_revenue'])
        self.assertEqual(self.engine.dimension_text('salesperson_id',MISSING),'未归属')

    def test_fine_and_coarse_budget_and_empty_object_selection(self):
        selected=Filters(dimensions={'salesperson_id':['1']})
        fine=sales_budget_comparison(self.engine,'2026-09',filters=selected,version='B2026_TEAM')
        coarse=sales_budget_comparison(self.engine,'2026-09',filters=selected,version='B2026_COMPANY')
        self.assertEqual(fine['status'],'complete')
        self.assertEqual(fine['totals']['net_revenue']['budget'],Decimal(4700))
        self.assertEqual(coarse['status'],'unsupported_dimensions')
        self.assertEqual(coarse['totals']['net_revenue']['actual'],Decimal(4300))
        for field in ('budget','completion_rate','variance'):
            self.assertIsNone(coarse['totals']['net_revenue'][field])
        all_objects=Filters(dimensions={'salesperson_id':None})
        self.assertTrue(all_objects.is_company)
        self.assertEqual(sales_budget_comparison(self.engine,'2026-09',filters=all_objects,version='B2026_COMPANY')['status'],'complete')

    def test_invalid_dimensions_are_never_ignored(self):
        for action in (lambda:self.engine.summary('2026-09',filters=Filters(dimensions={'wrong':['1']})),lambda:self.engine.dimensions('2026-09',by=['salesperson_id','salesperson_id'])):
            with self.assertRaises(ValueError):
                action()

    def test_model_grain_fixed_and_insufficient_history_not_company_forecast(self):
        training=History(self.engine).snapshot('2026-09')
        period=predict(training,['2026-10'],'recent_mean',3)[0]
        self.assertIn('salesperson_id',period['unavailable_dimensions'])
        selected=period_summary(period,training.flags,Filters(dimensions={'salesperson_id':['1']}))
        self.assertFalse(selected['is_available'])
        self.assertIsNone(selected['net_revenue'])
        self.assertEqual(selected['reason'],'dimension_history_incomplete')
        company=period_summary(period,training.flags)
        self.assertIsNotNone(company['net_revenue'])
        store=period_summary(period,training.flags,Filters(dimensions={'store_id':['SHOP2']}))
        self.assertIsNotNone(store['net_revenue'])
        excluded=period_summary(period,training.flags,Filters(dimensions={'customer_type':['企业客户']}))
        self.assertIsNone(excluded['net_revenue'])
        self.assertNotIn('customer_type',training.dimensions)

    def test_coarser_model_does_not_invent_or_discard_available_company_budget(self):
        changed=deepcopy(self.facts)
        missing=changed.tables['sales']['salesperson_id'].eq(MISSING)
        changed.tables['sales'].loc[missing,'salesperson_id']='6'
        changed.tables['sales'].loc[missing,'sales_department_id']='D2'
        changed.metadata['forecast_dimensions']=list(BUILTINS)
        engine=FinancialEngine(changed)
        base=forecast_base(engine,'2026-06','recent_mean',3,'B2026_TEAM')
        pack=forecast_pack(base)
        self.assertIsNotNone(pack['budget_reference']['company']['net_revenue'])
        fine=forecast_pack(base,Filters(dimensions={'salesperson_id':['1']}))
        self.assertIsNone(fine['scenarios']['base']['selected_remaining']['net_revenue'])
        self.assertIsNotNone(fine['budget_reference']['selected']['net_revenue'])

    def test_cutoff_rejects_future_classification_leakage(self):
        changed=deepcopy(self.facts)
        future=changed.tables['sales']['revenue_date']>date(2026,6,30)
        changed.tables['sales'].loc[future,'salesperson_id']='FUTURE'
        original=History(self.engine).snapshot('2026-06')
        altered=History(FinancialEngine(changed)).snapshot('2026-06')
        self.assertEqual(original.fingerprint,altered.fingerprint)
        self.assertEqual(predict(original,['2026-07'],'recent_mean',3),predict(altered,['2026-07'],'recent_mean',3))

    def test_forecast_group_does_not_sum_missing_cost_as_zero(self):
        import pandas as pd
        rows=[{'month':date(2026,10,1),'sales_department_id':dept,'sales_revenue':Decimal(10),'sales_cost':cost,'refund_amount':Decimal(1),'cogs_reversal_amount':Decimal(0)} for dept,cost in [('A',Decimal(3)),('A',None),('B',Decimal(3))]]
        grouped=forecast_groups(pd.DataFrame(rows,dtype=object),['sales_department_id']).set_index('sales_department_id')
        self.assertEqual(grouped.loc['A','net_revenue'],Decimal(18))
        self.assertIsNone(grouped.loc['A','net_cost'])
        self.assertIsNone(grouped.loc['A','gross_margin'])
        self.assertIsNone(grouped.loc['A','gross_margin_rate'])
        self.assertEqual(grouped.loc['B','gross_margin'],Decimal(6))
        self.assertEqual(grouped.loc['B','gross_margin_rate'],Decimal(6)/Decimal(9))

    def test_independent_legacy_baseline_unchanged(self):
        engine=FinancialEngine(load_facts(ROOT/'configs/demo_files.json'))
        result=engine.summary('2026-09')
        self.assertEqual([result[k] for k in ('net_revenue','net_cost','gross_margin','expenses','operating_profit')],list(map(Decimal,['267927.12','165822.96','102104.16','60504.74','41599.42'])))

    def test_report_and_excel_keep_dimensions_filters_and_complete_groups(self):
        selected=Filters(dimensions={'store_id':['SHOP2']})
        groups=['sales_department_id','salesperson_id']
        report=build_report(self.engine,'2026-09',selected,'B2026_COMPANY',method='recent_mean',group_by=groups)
        self.assertEqual(report['metadata']['group_by'],groups)
        self.assertEqual(report['metadata']['sales_filters'],selected.as_dict())
        self.assertIn('门店',report['metadata']['scope_label'])
        self.assertEqual(sum(row['net_revenue'] for row in report['sources']['grouped']),report['sources']['analysis']['selected_scope']['net_revenue'])
        self.assertIn('salesperson_id',report['sources']['details']['sales'][0])
        book=load_workbook(BytesIO(excel_bytes(report)))
        self.assertEqual(len(book.sheetnames),10)
        self.assertIn('联合分组；跨组订单数不可直接相加',[cell.value for row in book['维度分析'] for cell in row])
        self.assertIn('销售员',[cell.value for row in book['销售明细'] for cell in row])
        other=build_report(self.engine,'2026-09',selected,'B2026_COMPANY',method='recent_mean',group_by=groups[::-1])
        self.assertNotEqual(report['report_id'],other['report_id'])
        self.assertNotEqual(safe_export_name('E','2026-09','销售',selected),safe_export_name('E','2026-09','销售',Filters(dimensions={'store_id':['SHOP1']})))


class DimensionImportTests(unittest.TestCase):
    def test_scan_numeric_employee_manual_fields_and_no_money_recommendation(self):
        config=profile()
        config['dimensions']=[]
        config.pop('forecast_dimensions')
        scan=inspect_profile(config)
        employees=[x for x in scan['candidates'] if x['field']=='employee_id']
        self.assertEqual(len(employees),1)
        self.assertTrue(employees[0]['recommended'])
        self.assertTrue(all('amount' not in item['field'] for item in scan['candidates']))
        self.assertIn('type',employees[0])
        self.assertIn('employee_name',scan['fields']['employees'])
        raw={'sales':[{'员工编码':101,'自定义类型':'A','销售金额':100,'日期':'2026-01-01'}]}
        custom=recommend({'tables':{'sales':{}},'dimensions':[]},raw)
        self.assertTrue(any(item['field']=='员工编码' and item['recommended'] for item in custom))
        self.assertTrue(any(item['field']=='自定义类型' and not item['recommended'] for item in custom))

    def test_many_to_one_names_and_audited_totals(self):
        facts=load_facts(profile())
        audit=facts.metadata['join_validation']
        self.assertEqual(audit['sales_rows_before'],audit['sales_rows_after'])
        self.assertTrue(audit['amounts_unchanged'])
        self.assertEqual(FinancialEngine(facts).dimension_text('salesperson_id','1'),'林悦 · 1')

    def test_lookup_duplicate_blocks_instead_of_multiplying_revenue(self):
        config=profile()
        from etl.sources import read_source_tables, normalize_mapped
        raw={}
        records,_=read_source_tables(config,raw_records=raw)
        raw['employees'].append(deepcopy(raw['employees'][0]))
        with self.assertRaisesRegex(ValueError,'唯一'):
            attach_dimensions(config,raw,normalize_mapped(records,config['entity_id']))

    def test_invalid_mapping_and_financial_field_collision_rejected(self):
        for identifier in ('sales_revenue','net_revenue','department','refund_amount','product_id'):
            config=profile()
            config['dimensions'][0]['id']=identifier
            with self.assertRaises(ValueError):
                validate_dimensions(config)
        config=profile()
        config['dimensions'][0]['bindings']['sales']['field']='NONEXISTENT'
        with self.assertRaisesRegex(ValueError,'不存在'):
            load_facts(config)

    def test_empty_csv_preserves_headers_for_manual_mapping(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'empty.csv'
            path.write_text('sale_line_id,employee_id\n',encoding='utf-8')
            config=profile()
            config['tables']['sales']['files']=[str(path)]
            self.assertIn('employee_id',inspect_profile(config)['fields']['sales'])

    def test_v1_remains_supported_and_credentials_are_rejected(self):
        old=json.loads((ROOT/'configs/sample_files.json').read_text(encoding='utf-8'))
        validate_profile(old)
        self.assertEqual(len(FinancialEngine(load_facts(old)).dimension_ids),3)
        old['source']['connection']={'password':'never'}
        with self.assertRaises(ValueError):
            validate_profile(old)


class DimensionMemoryTests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.path=Path(self.folder.name)/'runtime/views.json'
        self.store=DimensionStore(self.path)
        self.known=[*BUILTINS,'salesperson_id','sales_department_id']
        self.layout={'slots':['product_id','channel_id','salesperson_id','region_id'],'group_by':['sales_department_id','salesperson_id']}

    def test_round_trip_without_objects_month_credentials_or_assumptions(self):
        self.store.remember('E:A',{**self.layout,'password':'forbidden','objects':['secret'],'month':'2026-09'},self.known)
        restored,unknown,active=DimensionStore(self.path).layout('E:A',self.known)
        self.assertEqual(restored,self.layout)
        self.assertEqual(unknown,[])
        self.assertIsNone(active)
        text=self.path.read_text(encoding='utf-8')
        for forbidden in ('password','secret','2026-09','objects'):
            self.assertNotIn(forbidden,text)
        self.assertEqual(self.store.layout('OTHER:A',self.known)[0]['slots'],list(BUILTINS))

    def test_recent_layout_does_not_overwrite_named_scheme(self):
        identity=self.store.save_scheme('E:A','销售团队',self.layout,self.known)
        changed=deepcopy(self.layout)
        changed['slots'][0]=None
        self.store.remember('E:A',changed,self.known,identity)
        self.assertEqual(self.store.schemes('E:A')[0]['layout'],self.layout)
        self.store.save_scheme('E:A','改名',changed,self.known,identity)
        self.assertEqual(self.store.schemes('E:A')[0]['layout'],changed)
        self.assertEqual(self.store.schemes('E:A')[0]['name'],'改名')
        self.store.delete('E:A',identity)
        self.assertEqual(self.store.schemes('E:A'),[])

    def test_scheme_export_import_checked_all_before_save(self):
        self.store.save_scheme('A','销售团队',self.layout,self.known)
        payload=self.store.export('A',[{'id':key,'label':key} for key in self.known])
        self.assertEqual(self.store.import_schemes('B',payload,self.known),1)
        before=self.path.read_bytes()
        with self.assertRaises(ValueError):
            self.store.import_schemes('B',payload,self.known)
        with self.assertRaises(ValueError):
            self.store.import_schemes('C',payload,list(BUILTINS))
        self.assertEqual(self.path.read_bytes(),before)

    def test_renaming_stable_id_restores_deleted_id_warns(self):
        self.store.remember('A',self.layout,self.known)
        restored,unknown,_=self.store.layout('A',[*BUILTINS,'sales_department_id'])
        self.assertEqual(unknown,['salesperson_id'])
        self.assertIsNone(restored['slots'][2])
        self.assertNotIn('salesperson_id',restored['group_by'])
        with self.assertRaises(ValueError):
            clean_layout({'slots':['product_id','product_id']},self.known)

    def test_stable_context_across_refresh_and_confirmed_mapping(self):
        config=profile()
        identity=profile_key(config)
        changed=deepcopy(config)
        changed['coverage']['complete_end']='2026-10-31'
        changed['dimensions'][0]['label']='改显示名'
        self.assertEqual(profile_key(changed),identity)
        self.store.confirm(config,config['dimensions'])
        restored=self.store.confirmed_profile(changed)
        self.assertEqual(restored['forecast_dimensions'],config['forecast_dimensions'])
        self.assertEqual(restored['coverage'],changed['coverage'])

    def test_updated_source_dimension_deletion_does_not_resurrect_saved_mapping(self):
        original=profile()
        confirmed=deepcopy(original)
        confirmed['dimensions'][1]['label']='客户经理'
        self.store.confirm(confirmed,confirmed['dimensions'],base_profile=original)
        self.assertEqual(self.store.confirmed_profile(original)['dimensions'][1]['label'],'客户经理')
        modified=deepcopy(original)
        modified['dimensions']=[item for item in modified['dimensions'] if item['id']!='salesperson_id']
        modified['forecast_dimensions'].remove('salesperson_id')
        restored=self.store.confirmed_profile(modified)
        self.assertNotIn('salesperson_id',[item['id'] for item in restored['dimensions']])
        renamed=deepcopy(original)
        renamed['dimensions'][1]['label']='新来源人员名'
        self.assertEqual(self.store.confirmed_profile(renamed)['dimensions'][1]['label'],'新来源人员名')

    def test_corrupt_and_write_failure_do_not_destroy_original(self):
        self.path.parent.mkdir(parents=True)
        self.path.write_text('broken',encoding='utf-8')
        self.assertFalse(self.store.remember('A',self.layout,self.known))
        self.assertEqual(self.path.read_text(encoding='utf-8'),'broken')
        self.path.unlink()
        store=DimensionStore(self.path)
        with patch('dashboard.dimension_store.os.replace',side_effect=OSError('readonly')):
            self.assertFalse(store.remember('A',self.layout,self.known))
        self.assertIn('本次',store.warning)
        self.assertEqual(list(self.path.parent.glob('.dimensions-*')),[])


if __name__=='__main__':
    unittest.main()
