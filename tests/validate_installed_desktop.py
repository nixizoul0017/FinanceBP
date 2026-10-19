"""Exercise installed resources in desktop mode, separate from native UI checks."""
import argparse
import json
import os
from pathlib import Path
import sys
import time
from unittest.mock import patch


def run(resources,data):
    resources,data=resources.resolve(),data.resolve()
    data.mkdir(parents=True,exist_ok=True)
    session=data/'runtime/sessions/acceptance'
    session.mkdir(parents=True,exist_ok=True)
    os.environ.update(FINANCEBP_DESKTOP='1',FINANCEBP_DATA_DIR=str(data),FINANCEBP_RESOURCES_DIR=str(resources),FINANCEBP_RUN_DIR=str(session))
    sys.path.insert(0,str(resources/'src'))
    from streamlit.testing.v1 import AppTest
    from desktop.artifacts import read_request,save_request,finish_request
    from desktop.events import atomic_json
    from desktop.host import Host
    from dashboard.uploads import SelectedFile
    from dashboard.navigation import PAGES
    from dashboard.dimension_store import DimensionStore
    from analysis import Filters
    from finance.budgets import workbook_bytes
    from openpyxl import load_workbook
    from io import BytesIO
    checks=[]
    def check(name,condition):
        checks.append({'check':name,'passed':bool(condition)})
        if not condition: raise AssertionError(name)
    def clean(app):
        check('render_'+app.session_state['page']+'_'+app.session_state['ui_theme'],not app.exception and not app.error)
    started=time.perf_counter()
    app=AppTest.from_file(resources/'app.py',default_timeout=120)
    app.session_state['source']='多维模拟数据'
    app.run();clean(app)
    engine=app.session_state['engine']
    check('installed_multi_salesperson_4300',engine.summary('2026-09',filters=Filters(dimensions={'salesperson_id':['1']}))['net_revenue']==4300)
    for theme in ('navy','teal','plum'):
        app.radio('ui_theme').set_value(theme).run()
        for page in PAGES:
            app.button('nav_'+page).click().run();clean(app)
    app.button('nav_预算工作台').click().run()
    next(b for b in app.button if b.label=='下载当前企业预算模板').click().run();clean(app)
    requests=list((session/'exports').glob('*.json'))
    check('native_budget_export_queued_once',len(requests)==1)
    request,payload=read_request(requests[0])
    book=load_workbook(BytesIO(payload));check('native_export_is_budget_workbook','销售预算' in book.sheetnames);book.close()
    target=save_request(requests[0],data/'中文预算模板.xlsx')
    check('native_save_unicode_path',target.exists())
    # Exercise the real host file-selection validation using a controlled chosen
    # file. The Windows dialog itself is checked separately with Computer Use.
    from unittest.mock import MagicMock
    host=Host.__new__(Host)
    host.window=MagicMock()
    host.window.create_file_dialog.return_value=(str(target),)
    upload_id='abcdef0123456789abcdef0123456789'
    upload=session/'uploads'/(upload_id+'.json')
    atomic_json(upload,{'id':upload_id,'types':['xlsx'],'limit':20*1024*1024})
    host.upload(upload)
    result=json.loads(upload.with_suffix('.result').read_text(encoding='utf-8'))
    check('host_upload_reads_unicode_file',result['state']=='selected' and upload.with_suffix('.bin').read_bytes()==payload)
    # Scheme upload keeps its stable identity and validates dimensions.
    store=DimensionStore()
    engine=app.session_state['engine']
    from etl.dimensions import profile_key
    profile=json.loads((resources/'configs/multidimensional_files.json').read_text(encoding='utf-8'))
    layout={'slots':['sales_department_id','salesperson_id','region_id','product_id'],'group_by':['sales_department_id','salesperson_id']}
    identity=store.save_scheme(profile_key(profile),'安装验收方案',layout,engine.dimension_ids)
    check('scheme_fourth_slot_and_joint_groups',DimensionStore().schemes(profile_key(profile))[0]['layout']==layout and bool(identity))
    return {'passed':True,'kind':'installed_resources_AppTest_with_locked_build_runtime',
            'native_dialog_manual_test':False,'resources':str(resources),'checks':checks,'seconds':round(time.perf_counter()-started,3)}


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--resources',type=Path,required=True)
    parser.add_argument('--data',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    try: result=run(args.resources,args.data)
    except Exception as error:
        import traceback
        result={'passed':False,'error':str(error),'traceback':traceback.format_exc()}
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({key:result[key] for key in ('passed','seconds','error') if key in result}))
    raise SystemExit(0 if result['passed'] else 1)
