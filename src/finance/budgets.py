"""Portable budget templates and immutable published overlays; never write source SQL."""
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from analysis.forecast_store import content_digest
from analysis.metrics import FinancialEngine, month_date
from etl.common import ZERO
from etl.dimensions import MISSING
from etl.sources import CONTRACT, Facts
from .store import WorkspaceError

FORMAT='financebp_budget_v1'
FEE_KEYS=('expense_category','department','channel_scope','region_scope')
LABELS={'month':'月份','budget_net_revenue':'预算净收入（元）','budget_net_cost':'预算净成本（元，可空）',
        'quantity':'成交数量（件）','net_unit_price':'退款前均价（元）','unit_cost':'单位历史成本（元，可空）',
        'refund_amount':'退款额（元）','cost_reversal_amount':'成本转回（元）','amount':'公司费用（元）',
        'expense_category':'费用类别','department':'费用部门','channel_scope':'费用渠道','region_scope':'费用地区'}


def number(value,field,optional=False):
    if value is None or (isinstance(value,str) and not value.strip()):
        if optional:
            return None
        raise WorkspaceError(LABELS.get(field,field)+'不能为空；已知为零请填0')
    if isinstance(value,bool):
        raise WorkspaceError('金额和数量不能使用布尔值')
    try:
        result=Decimal(str(value))
    except (InvalidOperation,ValueError) as error:
        raise WorkspaceError(LABELS.get(field,field)+'必须是有效数字') from error
    if not result.is_finite() or result<0 or result>Decimal('1000000000000'):
        raise WorkspaceError(LABELS.get(field,field)+'需为0至1万亿元内的有限数字')
    places=0 if field=='quantity' else 6 if field in ('unit_cost','net_unit_price') else 2
    if result!=result.quantize(Decimal(1).scaleb(-places)):
        raise WorkspaceError(LABELS.get(field,field)+f'最多允许{places}位小数')
    return int(result) if field=='quantity' else result


def normalize_dataset(dataset):
    d=deepcopy(dataset)
    if not isinstance(d,dict) or d.get('format')!=FORMAT:
        raise WorkspaceError('预算模板格式或版本不受支持')
    if not isinstance(d.get('entity_id'),str) or not d['entity_id'].strip():
        raise WorkspaceError('预算缺少企业编号')
    if type(d.get('year')) is not int or not 2000<=d['year']<=2100:
        raise WorkspaceError('预算年度需在2000—2100之间')
    if d.get('input_mode') not in ('amounts','drivers'):
        raise WorkspaceError('预算输入方式不受支持')
    for key in ('grain','expense_grain'):
        fields=d.get(key)
        if not isinstance(fields,list) or len(set(fields))!=len(fields) or any(not isinstance(x,str) or not x for x in fields):
            raise WorkspaceError('预算分类配置无效')
    if not set(d['expense_grain']).issubset(FEE_KEYS):
        raise WorkspaceError('费用预算分类无效')
    fields=['budget_net_revenue','budget_net_cost'] if d['input_mode']=='amounts' else ['quantity','net_unit_price','unit_cost','refund_amount','cost_reversal_amount']
    for kind,grain,amounts in [('sales',d['grain'],fields),('expenses',d['expense_grain'],['amount'])]:
        if not isinstance(d.get(kind),list) or len(d[kind])>100000:
            raise WorkspaceError('预算明细必须为最多10万行的列表')
        seen=set()
        for row in d[kind]:
            if not isinstance(row,dict) or set(row)!=set(['month',*grain,*amounts]):
                raise WorkspaceError('预算明细字段与模板分类配置不一致')
            raw=row['month']
            if isinstance(raw,datetime):
                raw=raw.date()
            try:
                month=month_date(raw)
            except (ValueError,TypeError) as error:
                raise WorkspaceError('月份必须为YYYY-MM或每月1日') from error
            if month.year!=d['year']:
                raise WorkspaceError('明细月份不属于模板预算年度')
            row['month']=month.isoformat()
            for field in grain:
                value=row[field]
                if value is None or not str(value).strip() or len(str(value))>200:
                    raise WorkspaceError('分类编码不能为空；未归属需使用明确编码')
                row[field]=str(value).strip()
            key=(row['month'],*(row[field] for field in grain))
            if key in seen:
                raise WorkspaceError('同月同分类存在重复预算行：'+str(key))
            seen.add(key)
            for field in amounts:
                row[field]=number(row.get(field),field,field in ('budget_net_cost','unit_cost'))
            if d['input_mode']=='drivers' and kind=='sales':
                if row['quantity']==0 and (row['refund_amount'] or row['cost_reversal_amount']):
                    # Cross-period refunds are allowed, including a refund-only budget month.
                    pass
    if not d['sales'] and not d['expenses']:
        raise WorkspaceError('没有已填写的预算；请至少填写一条销售或费用预算')
    d['warnings']=list(d.get('warnings',[]))
    return d


def validate_dataset(dataset):
    return normalize_dataset(dataset)


def check_source(dataset,engine,allow_new=False):
    d=normalize_dataset(dataset)
    if d['entity_id']!=engine.entity_id:
        raise WorkspaceError('模板企业与当前企业不一致，请重新下载当前企业模板')
    if not set(d['grain']).issubset(engine.dimension_ids):
        raise WorkspaceError('模板含未确认或已失效的分类，请核对接入配置')
    warnings=[]
    for field in d['grain']:
        known=set(engine.tables['sales'][field])
        known.update(engine.tables['sales_budget'][field] if field in engine.tables['sales_budget'] else [])
        unknown=sorted({row[field] for row in d['sales']}-known)
        if unknown:
            warnings.append(f'{engine.dimension_labels[field]}新增计划编码：'+ '、'.join(unknown[:10]))
    for field in d['expense_grain']:
        known=set(engine.tables['expenses'][field])|set(engine.tables['expense_budget'][field])
        unknown=sorted({row[field] for row in d['expenses']}-known)
        if unknown:
            warnings.append(LABELS[field]+'新增计划编码：'+'、'.join(unknown[:10]))
    if warnings and not allow_new:
        raise WorkspaceError('；'.join(warnings)+'。如为明确的新业务计划，请勾选允许新增计划编码后重新校验')
    d['warnings']=warnings
    return d


def columns(d,kind):
    return ['month',*(d['grain'] if kind=='sales' else d['expense_grain']),
            *(['budget_net_revenue','budget_net_cost'] if d['input_mode']=='amounts' else ['quantity','net_unit_price','unit_cost','refund_amount','cost_reversal_amount'])] if kind=='sales' else ['month',*d['expense_grain'],'amount']


def workbook_bytes(engine,year,grain=(),input_mode='amounts',expense_grain=(),dataset=None):
    d=normalize_dataset(dataset) if dataset else {'format':FORMAT,'entity_id':engine.entity_id,'year':int(year),'grain':list(grain),'input_mode':input_mode,'expense_grain':list(expense_grain),'sales':[],'expenses':[]}
    if not set(d['grain']).issubset(engine.dimension_ids) or not set(d['expense_grain']).issubset(FEE_KEYS) or d['input_mode'] not in ('amounts','drivers'):
        raise WorkspaceError('模板分类或输入方式无效')
    wb=Workbook()
    info=wb.active
    info.title='说明'
    info.append(['FinanceBP · 预算编制模板','填写前请阅读'])
    for field in ('format','entity_id','year','input_mode','grain','expense_grain'):
        value=d[field]
        info.append([field,json.dumps(value,ensure_ascii=False) if isinstance(value,list) else value])
    notes=[('填写方式','保留字段编号行；月份填写每月1日或YYYY-MM，分类填写编码，不能填写显示名称。'),
           ('金额口径','人民币、不含税、元；净收入/净成本均为扣除退款/成本转回后的金额，不填含税数。'),
           ('空值规则','空白成本代表未知，不按零处理；已知零预算请明确填0。只有月份的未填写行会忽略。'),
           ('费用范围','费用始终为公司全口径，费用部门不会自动对应销售部门。'),
           ('版本流程','上传校验 → 保存草稿 → 确认发布；发布后不可覆盖，调整时另建关联版本。'),
           ('公式规则','可填写数值或带Excel已计算缓存的内部公式；无缓存、外部引用公式会被拒绝。'),
           ('安全说明','本模板只维护程序本地预算，不修改企业数据库。说明表前7行请勿修改。')]
    for row in notes:
        info.append(row)
    reference=wb.create_sheet('分类参考')
    reference.append(['稳定编号','显示名称','分类编码','对象名称'])
    for entry in engine.registry:
        field=entry['id']
        values=sorted(set(engine.tables['sales'][field]))
        for value in values[:10000]:
            reference.append([field,entry['label'],value,entry.get('names',{}).get(value,value)])
    for name,kind in [('销售预算','sales'),('费用预算','expenses')]:
        ws=wb.create_sheet(name)
        fields=columns(d,kind)
        ws.append([LABELS.get(field,engine.dimension_labels.get(field,field)) for field in fields])
        ws.append(fields)
        for record in d[kind]:
            ws.append([date.fromisoformat(record[field]) if field=='month' else float(record[field]) if isinstance(record[field],Decimal) else record[field] for field in fields])
        if not d[kind]:
            for m in range(1,13):
                ws.append([date(d['year'],m,1),*[None]*(len(fields)-1)])
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.font=Font(name='Calibri',size=11,color='2161B0')
                cell.number_format='yyyy-mm' if cell.column==1 else '#,##0.00;[Red](#,##0.00);"—"' if fields[cell.column-1] in LABELS and fields[cell.column-1] not in FEE_KEYS else '@'
        for index,field in enumerate(fields,1):
            if field in d['grain']:
                refs=[r for r in range(2,reference.max_row+1) if reference.cell(r,1).value==field]
                if refs:
                    from openpyxl.utils import get_column_letter
                    dv=DataValidation(type='list',formula1=f"'分类参考'!$C${min(refs)}:$C${max(refs)}",allow_blank=True)
                    dv.errorTitle='请填写编码'
                    dv.promptTitle='分类编码'
                    dv.prompt='可从分类参考中查找；新计划编码需在导入时确认。'
                    dv.showInputMessage=True
                    ws.add_data_validation(dv)
                    dv.add(f'{get_column_letter(index)}3:{get_column_letter(index)}100002')
        ws.freeze_panes='B3'
        ws.auto_filter.ref=f'A2:{ws.cell(ws.max_row,ws.max_column).coordinate}'
    # Order instructions → inputs → reference.
    wb._sheets=[info,wb['销售预算'],wb['费用预算'],reference]
    for ws in wb:
        ws.sheet_view.showGridLines=False
        ws.freeze_panes=ws.freeze_panes or 'A2'
        ws.row_dimensions[1].height=32
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='19374D')
            cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=12)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=25 if ws.title!='说明' else 28 if col[0].column==1 else 95
        for row in ws:
            if row[0].row>1:
                ws.row_dimensions[row[0].row].height=30 if ws.title=='说明' else 24
            for cell in row:
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if cell.row>1 and cell.data_type=='f':
                    raise WorkspaceError('分类文本含公式内容')
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):
                    cell.data_type='s'
    output=BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_workbook(payload,engine,allow_new=False):
    if not isinstance(payload,bytes) or len(payload)>20*1024*1024:
        raise WorkspaceError('请上传不超过20MB的xlsx预算模板')
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(entry.file_size for entry in archive.infolist())>100*1024*1024:
                raise WorkspaceError('工作簿展开后过大，请拆分导入')
        wb=load_workbook(BytesIO(payload),data_only=False,read_only=True,keep_links=False)
        cached=load_workbook(BytesIO(payload),data_only=True,read_only=True,keep_links=False)
        try:
            if not {'说明','销售预算','费用预算'}.issubset(wb.sheetnames):
                raise WorkspaceError('请使用本程序下载的预算模板')
            meta={row[0]:row[1] for row in wb['说明'].iter_rows(min_row=2,max_row=7,max_col=2,values_only=True)}
            d={**meta,'year':int(meta['year']),'grain':json.loads(meta['grain']),'expense_grain':json.loads(meta['expense_grain']),'sales':[],'expenses':[]}
            errors=[]
            for sheet,kind in [('销售预算','sales'),('费用预算','expenses')]:
                fields=columns(d,kind)
                ws=wb[sheet]
                if ws.max_row>100002 or ws.max_column>len(fields):
                    raise WorkspaceError(sheet+'超出模板范围')
                if [c.value for c in next(ws.iter_rows(min_row=2,max_row=2))]!=fields:
                    raise WorkspaceError(sheet+'的字段编号行已变化，请重新下载模板')
                cache_rows=cached[sheet].iter_rows(min_row=3,values_only=True)
                for index,(cells,values) in enumerate(zip(ws.iter_rows(min_row=3),cache_rows),3):
                    actual=[]
                    for cell,value in zip(cells,values):
                        if cell.data_type=='f' and (value is None or '[' in cell.value or '!' in cell.value):
                            errors.append(f'{sheet}第{index}行公式没有可用缓存或引用其他表，请在Excel中粘贴为数值')
                        if cell.data_type=='e':
                            errors.append(f'{sheet}第{index}行含Excel错误值')
                        actual.append(value if cell.data_type=='f' else cell.value)
                    if all(value is None or value=='' for value in actual[1:]):
                        continue
                    d[kind].append(dict(zip(fields,actual)))
            if errors:
                raise WorkspaceError('；'.join(errors[:10]))
            return check_source(d,engine,allow_new)
        finally:
            wb.close()
            cached.close()
    except WorkspaceError:
        raise
    except Exception as error:
        raise WorkspaceError('无法识别预算工作簿，请使用当前企业的xlsx模板，保留说明和字段编号行') from error


def overlay(engine,versions):
    if not versions:
        return engine
    tables={key:frame.copy(deep=True) for key,frame in engine.tables.items()}
    meta=deepcopy(engine.metadata)
    meta.setdefault('budget_dimensions',{})
    meta.setdefault('expense_budget_dimensions',{})
    meta.setdefault('budget_input_modes',{})
    meta.setdefault('budget_labels',{})
    manifest=[]
    for version in versions:
        if version['status']!='published':
            continue
        d=check_source(version['dataset'],engine,allow_new=True)
        identity=version['id']
        if identity in set(tables['sales_budget']['budget_version'])|set(tables['expense_budget']['budget_version']):
            raise WorkspaceError('本地预算编号与来源版本冲突')
        meta['budget_dimensions'][identity]=d['grain']
        meta['expense_budget_dimensions'][identity]=d['expense_grain']
        meta['budget_input_modes'][identity]=d['input_mode']
        meta['budget_labels'][identity]=version['name']+' · 本机'+('调整预算' if version['kind']=='adjustment' else '原始预算')
        for kind,name,grain in [('sales','sales_budget',d['grain']),('expenses','expense_budget',d['expense_grain'])]:
            records=[]
            for row in d[kind]:
                value={**row,'entity_id':engine.entity_id,'budget_version':identity,'month':date.fromisoformat(row['month'])}
                if kind=='sales':
                    value.update({field:value.get(field,MISSING) for field in engine.dimension_ids})
                    if d['input_mode']=='amounts':
                        value.update(quantity=0,net_unit_price=ZERO,unit_cost=None,refund_amount=ZERO,cost_reversal_amount=ZERO,budget_input_kind='amounts')
                else:
                    value.update({field:value.get(field,'ALL') for field in FEE_KEYS})
                records.append(value)
            if records:
                tables[name]=pd.concat([tables[name],pd.DataFrame(records,dtype=object)],ignore_index=True)
                meta.setdefault('data_availability',{})[name]=True
        manifest.append({field:version[field] for field in ('id','name','kind','parent_id','digest','approver','updated_at')})
    meta['local_budgets']=manifest
    meta['actual_source_fingerprint']=engine.metadata.get('actual_source_fingerprint',engine.metadata.get('source_fingerprint'))
    meta['source_fingerprint']=content_digest({'source':engine.metadata.get('source_fingerprint'),'local_budgets':manifest})
    return FinancialEngine(Facts(tables,meta))
