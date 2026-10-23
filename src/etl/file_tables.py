"""Read local business tables without executing workbook macros or changing codes."""
from dataclasses import dataclass
from datetime import date, datetime
import codecs
import csv
import json
from pathlib import Path
import re
import zipfile

import pandas as pd

FILE_TYPES = ('xlsx', 'xls', 'xlsm', 'dta', 'csv', 'tsv', 'json')
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_ROWS = 1_000_000
MAX_CELLS = 5_000_000


@dataclass
class TableData:
    columns: list
    rows: list
    comments: dict


def _path(value):
    path = Path(value)
    if path.suffix.lower().lstrip('.') not in FILE_TYPES:
        raise ValueError('支持Excel（xlsx/xls/xlsm）、Stata（dta）、CSV、TSV及记录列表JSON。')
    if not path.is_file():
        raise ValueError('找不到业务文件，请重新选择文件。')
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError('单个业务文件不能超过100MB，请按业务表或期间拆分。')
    return path


def _xlsx_check(path):
    try:
        with zipfile.ZipFile(path) as book:
            if sum(info.file_size for info in book.infolist()) > 512 * 1024 * 1024:
                raise ValueError('Excel解压后的内容过大，请删除无关工作表、图片或拆分文件。')
    except zipfile.BadZipFile:
        raise ValueError('Excel文件内容与扩展名不符，或文件已加密/损坏。请另存为未加密的xlsx。') from None


def list_sheets(value):
    path = _path(value)
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        _xlsx_check(path)
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            return list(book.sheetnames)
        finally:
            book.close()
    if path.suffix.lower() == '.xls':
        try:
            import xlrd
        except ImportError:
            raise ValueError('缺少旧版Excel读取组件xlrd，请使用新版安装包或更新项目依赖。') from None
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    return [None]


def scalar(value):
    if value is None or (not isinstance(value, (str, bytes, list, dict)) and pd.isna(value)):
        return None
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if hasattr(value, 'item'):
        value = value.item()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (list, dict, tuple, bytes)):
        raise ValueError('业务数据须为二维明细表，单元格不能包含嵌套对象或二进制内容。')
    return value


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _records(header, iterator, limit, comments=None):
    header = [scalar(value) for value in header]
    while header and _blank(header[-1]):
        header.pop()
    if not header or any(_blank(value) for value in header):
        raise ValueError('表头有空白列，请选择正确的表头行，并为业务列填写名称。')
    columns = [str(value).strip().lstrip('\ufeff') for value in header]
    if len(columns) != len(set(columns)):
        raise ValueError('表头包含重复列名，请先改成不同名称，避免字段映射错误。')
    rows = []
    for values in iterator:
        values = [scalar(value) for value in values]
        if all(_blank(value) for value in values):
            continue
        if any(not _blank(value) for value in values[len(columns):]):
            raise ValueError('存在没有表头的数据列，请核对表头行和文件分隔符。')
        if len(rows) >= MAX_ROWS or (len(rows)+1)*len(columns) > MAX_CELLS:
            raise ValueError('业务表超过100万行或500万个单元格，请拆分文件。')
        rows.append(dict(zip(columns, values[:len(columns)] + [None]*max(0,len(columns)-len(values)))))
        if limit is not None and len(rows) >= limit:
            break
    return TableData(columns, rows, comments or {})


def _csv_encoding(path, encoding):
    if encoding != 'auto':
        return encoding
    with path.open('rb') as stream:
        prefix = stream.read(65536)
    if prefix.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return 'utf-16'
    try:
        codecs.getincrementaldecoder('utf-8-sig')().decode(prefix, final=False)
        return 'utf-8-sig'
    except UnicodeDecodeError:
        return 'gb18030'


def read_table(value, *, sheet_name=None, header_row=1, encoding='auto', limit=None):
    path = _path(value)
    if not isinstance(header_row, int) or not 1 <= header_row <= 100:
        raise ValueError('表头行应为1至100之间的整数。')
    suffix = path.suffix.lower()
    if suffix in ('.xlsx', '.xlsm'):
        _xlsx_check(path)
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if sheet_name is not None and sheet_name not in book.sheetnames:
                raise ValueError('找不到配置指定的Excel工作表。')
            sheet = book[sheet_name] if sheet_name is not None else book.worksheets[0]
            def values():
                for cells in sheet.iter_rows(min_row=header_row):
                    row = []
                    for cell in cells:
                        if cell.data_type == 'e':
                            raise ValueError('Excel包含错误单元格：'+cell.coordinate+'，请修正后重试。')
                        item = cell.value
                        if isinstance(item, (float,int)) and not isinstance(item,bool) and re.fullmatch('0{2,}',cell.number_format or '') and float(item).is_integer():
                            item = str(int(item)).zfill(len(cell.number_format))
                        row.append(item)
                    yield row
            iterator = iter(values())
            return _records(next(iterator, []), iterator, limit)
        finally:
            book.close()
    if suffix == '.xls':
        try:
            import xlrd
        except ImportError:
            raise ValueError('缺少旧版Excel读取组件xlrd，请使用新版安装包或更新项目依赖。') from None
        book = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
        try:
            sheet = book.sheet_by_name(sheet_name) if sheet_name is not None else book.sheet_by_index(0)
            def values():
                for index in range(header_row-1,sheet.nrows):
                    result = []
                    for cell in sheet.row(index):
                        if cell.ctype == xlrd.XL_CELL_ERROR:
                            raise ValueError('旧版Excel包含错误单元格，请修正后重试。')
                        value=xlrd.xldate_as_datetime(cell.value,book.datemode) if cell.ctype==xlrd.XL_CELL_DATE else cell.value
                        if cell.ctype==xlrd.XL_CELL_NUMBER:
                            format_code=book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                            if re.fullmatch('0{2,}',format_code) and float(value).is_integer():
                                value=str(int(value)).zfill(len(format_code))
                        result.append(value)
                    yield result
            iterator=iter(values())
            return _records(next(iterator,[]),iterator,limit)
        finally:
            book.release_resources()
    if suffix in ('.csv','.tsv'):
        selected = _csv_encoding(path,encoding)
        try:
            with path.open(encoding=selected,newline='') as stream:
                reader=csv.reader(stream,delimiter='\t' if suffix=='.tsv' else ',')
                for _ in range(header_row-1): next(reader,None)
                return _records(next(reader,[]),reader,limit)
        except UnicodeError:
            raise ValueError('文本文件编码无法读取，请另存为UTF-8 CSV/TSV，或在接入配置中指定编码。') from None
    if suffix == '.json':
        with path.open(encoding='utf-8-sig') as stream:
            rows=json.load(stream)
        if isinstance(rows,dict) and 'profile_version' in rows:
            raise ValueError('这是接入配置JSON，请切换到“已有JSON配置”入口。')
        if not isinstance(rows,list) or not all(isinstance(row,dict) for row in rows):
            raise ValueError('业务JSON应为记录对象列表；字段映射配置请使用“已有JSON配置”入口。')
        if not rows:
            return TableData([],[],{})
        columns=list(dict.fromkeys(key for row in rows for key in row))
        return _records(columns,([row.get(key) for key in columns] for row in rows),limit)
    if suffix == '.dta':
        # Keep numeric codes authoritative: labels are metadata, never substituted.
        try:
            with pd.read_stata(path,iterator=True,convert_categoricals=False,convert_dates=True) as reader:
                comments=reader.variable_labels()
                frame=reader.read(nrows=limit if limit is not None else MAX_ROWS+1)
            if frame.size > MAX_CELLS:
                raise ValueError('Stata表超过500万个单元格，请拆分后导入。')
            return _records(list(frame.columns),frame.itertuples(index=False,name=None),limit,comments)
        except (UnicodeError,NotImplementedError):
            raise ValueError('无法读取此Stata版本或编码，请在Stata中另存为支持Unicode的dta（118/119）。') from None
    raise ValueError('不支持的业务文件格式。')
