"""按原项目计划用openpyxl生成结果快照，计算只由财务模块负责。"""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import math
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from .tables import workbook_tables
from .service import validate_report

MONEY='#,##0.00;(#,##0.00);"-"'
RATIO='0.00%;(0.00%);"-"'
COUNT='#,##0;(#,##0);"-"'
DARK="203D52"


def write_cell(ws,row,column,value,kind):
    cell=ws.cell(row,column)
    if value is None:
        value="不可用"
    if isinstance(value,Decimal):
        if not value.is_finite():
            raise ValueError("不能导出非有限金额")
        value=float(value)
    if isinstance(value,float) and not math.isfinite(value):
        raise ValueError("不能导出非有限数字")
    if kind=="bool" and isinstance(value,bool):
        value="是" if value else "否"
    if isinstance(value,str):
        # 显式字符串类型避免企业编码/外部文本被解释为Excel公式。
        cell.value=re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]","",value)
        cell.data_type="s"
    else:
        cell.value=value
    cell.font=Font(name="Arial",size=11,color="203040")
    cell.alignment=Alignment(vertical="center",horizontal="right" if isinstance(value,(int,float,Decimal)) and not isinstance(value,bool) else "left")
    if isinstance(value,(date,datetime)):
        cell.number_format="yyyy-mm-dd"
    elif kind=="money":
        cell.number_format=MONEY
    elif kind=="ratio":
        cell.number_format=RATIO
    elif kind=="count":
        cell.number_format=COUNT
    elif kind=="metric":
        cell.number_format=RATIO if ws.cell(row,column-1).value=="毛利率" else MONEY
    elif kind=="number":
        cell.number_format="0.###############"
    elif kind=="check":
        cell.number_format='#,##0.00;(#,##0.00);0.00'
    return cell


def excel_bytes(report):
    validate_report(report)
    specs=workbook_tables(report)
    book=Workbook()
    book.remove(book.active)
    book.properties.title="经营分析月报"
    book.properties.subject=str(report["metadata"]["month"])[:7]+" "+report["metadata"]["entity_id"]
    book.properties.description=report["report_id"]+"；当次结果快照；精确金额见配套JSON"
    table_id=0
    for spec in specs:
        ws=book.create_sheet(spec["name"])
        ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.sheet_properties.outlinePr.summaryRight=False
        ws.sheet_properties.tabColor=DARK if spec["name"]=="总览点评" else "7893A6" if spec["name"]=="预测假设" else None
        ws.column_dimensions["A"].width=3
        ws["B2"]=spec["title"]
        ws["B2"].font=Font(name="Arial",size=16,bold=True,color=DARK)
        ws.row_dimensions[2].height=28
        context=[report["metadata"]["data_origin"]+"；人民币不含税元；预算 "+report["metadata"]["budget_version"],
            "销售："+report["metadata"]["scope_label"]+"；公司费用与利润保留公司全口径",
            "本次结果快照；调整假设后在网页重新生成。未知显示不可用，精确金额见配套JSON。"]
        widest=max(len(section["headers"]) for section in spec["sections"])
        for r,text in enumerate(context,3):
            cell=write_cell(ws,r,2,text,"text")
            if widest>1:
                ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=widest+1)
            cell.font=Font(name="Arial",size=11,italic=True,color="546575")
            cell.alignment=Alignment(wrap_text=True,vertical="center")
            ws.row_dimensions[r].height=30
        row=7
        # 内容列宽按所有区段取最大，文本区只对长文本列换行。
        widths={}
        for block in spec["sections"]:
            for index,kind in enumerate(block["types"],2):
                default=26 if kind=="text" else 23 if kind in ("money","metric") else 19
                width=block["widths"][index-2] if block["widths"] else default
                widths[index]=max(width,widths.get(index,0))
        # 总览点评专用版式：主指标与文字共享整页宽度，避免长点评撑开数字列。
        if spec["name"]=="总览点评":
            widths={2:24,3:26,4:25,5:25}
        elif spec["name"]=="预测假设":
            widths={column:24 for column in range(2,widest+2)}
        for column,width in widths.items():
            ws.column_dimensions[get_column_letter(column)].width=width
        first_header=None
        for block in spec["sections"]:
            write_cell(ws,row,2,block["title"],"text").font=Font(name="Arial",size=11,bold=True,color=DARK)
            ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=widest+1)
            ws.cell(row,2).alignment=Alignment(wrap_text=True,vertical="center")
            ws.row_dimensions[row].height=30
            row+=1
            if spec["name"]=="总览点评" and block["title"]=="经营点评":
                for values in block["rows"]:
                    text=values[0]+"（"+values[1]+"）："+values[2]+(" 依据："+values[3] if values[3] else "")
                    cell=write_cell(ws,row,2,text,"text")
                    ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=5)
                    cell.alignment=Alignment(wrap_text=True,vertical="center")
                    ws.row_dimensions[row].height=max(34,math.ceil(len(text)/66)*19)
                    row+=1
                row+=2
                continue
            header=row
            if first_header is None:
                first_header=header
            for col,label in enumerate(block["headers"],2):
                cell=write_cell(ws,row,col,label,"text")
                cell.fill=PatternFill("solid",fgColor=DARK)
                cell.font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
                cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
                cell.border=Border(bottom=Side(style="thin",color="DDE5EB"))
            ws.row_dimensions[row].height=34
            row+=1
            for values in block["rows"]:
                if len(values)!=len(block["headers"]):
                    raise ValueError("月报列数不一致")
                height=24
                for col,(value,kind) in enumerate(zip(values,block["types"]),2):
                    cell=write_cell(ws,row,col,value,kind)
                    if isinstance(value,str) and len(value)>widths[col]*1.5:
                        cell.alignment=Alignment(wrap_text=True,vertical="center")
                        height=max(height,min(180,math.ceil(len(value)/(widths[col]*0.85))*18+6))
                if spec["name"]=="预测假设" and block["title"]=="模型说明":
                    ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=widest+1)
                    ws.cell(row,3).alignment=Alignment(wrap_text=True,vertical="center")
                    height=max(26,math.ceil(len(str(values[1]))/90)*18+6)
                ws.row_dimensions[row].height=height
                row+=1
            if not block["rows"]:
                write_cell(ws,row,2,"此区段没有可用明细；请检查上方状态与口径说明。","text")
                ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=widest+1)
                ws.row_dimensions[row].height=28
                row+=1
            elif not (spec["name"]=="预测假设" and block["title"]=="模型说明"):
                table_id+=1
                table=Table(displayName="ReportTable"+str(table_id),ref=f"B{header}:{get_column_letter(len(block['headers'])+1)}{row-1}")
                table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False)
                ws.add_table(table)
            row+=2
        if spec["name"]!="总览点评" and first_header:
            ws.freeze_panes="C"+str(first_header+1)
        ws.print_options.horizontalCentered=True
        ws.page_setup.orientation="landscape"
        ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1
        ws.page_setup.fitToHeight=0
        ws.print_area=f"B2:{get_column_letter(widest+1)}{row-1}"
    output=BytesIO()
    book.save(output)
    return output.getvalue()


def inspect_excel(payload):
    """用实际导出文件读取类型和主要结果，供测试及验收。"""
    book=load_workbook(BytesIO(payload),data_only=False)
    errors=[]
    for ws in book:
        for row in ws:
            for cell in row:
                if cell.data_type in ("e","f"):
                    errors.append(ws.title+"!"+cell.coordinate)
    if errors:
        raise ValueError("快照含公式或错误单元格："+",".join(errors[:5]))
    return book
